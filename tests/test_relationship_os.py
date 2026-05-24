import importlib.util
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "relationship_os.py"


def load_relationship_os_module():
    spec = importlib.util.spec_from_file_location("relationship_os_under_test", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules["relationship_os_under_test"] = module
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


class RelationshipOsCliTest(unittest.TestCase):
    def run_cli(self, args, env):
        proc = subprocess.run(
            [sys.executable, str(SCRIPT), "--env", str(ROOT / ".env.example"), *args],
            cwd=ROOT,
            text=True,
            capture_output=True,
            env=env,
        )
        self.assertEqual(proc.returncode, 0, proc.stderr)
        return proc.stdout

    def sqlite_csv_env(self, base: Path):
        env = os.environ.copy()
        env["RELATIONSHIP_OS_STORE"] = "sqlite"
        env["RELATIONSHIP_OS_DB_PATH"] = str(base / "relationship_os.sqlite3")
        env["RELATIONSHIP_OS_CONSULTANT_VIEW"] = "csv"
        env["RELATIONSHIP_OS_LOCAL_DIR"] = str(base / "local_sheet")
        return env

    def sample_touchpoint_payload(self, contact_name="Demo Client", **overrides):
        payload = {
            "contact_name": contact_name,
            "touch_date": "2026-05-20",
            "touchpoint_type": "coffee",
            "sentiment": "positive",
            "summary": f"Coffee with {contact_name}. Interested in retirement planning.",
            "raw_input": f"Had coffee with {contact_name}. Follow up next Friday.",
            "topics": ["retirement"],
            "action_items": "Follow up next Friday on retirement planning interest.",
            "contact_type": "prospect",
            "relationship_stage": "warm",
        }
        payload.update(overrides)
        return payload

    def run_json_cli(self, args, env):
        return json.loads(self.run_cli(["--format=json", *args], env))

    def sqlite_row(self, db_path: Path, table: str, row_id: str):
        with sqlite3.connect(db_path) as conn:
            conn.row_factory = sqlite3.Row
            row = conn.execute(
                f'SELECT * FROM "{table}" WHERE id = ? LIMIT 1',
                (row_id,),
            ).fetchone()
        return dict(row) if row else None

    def test_local_demo_flow(self):
        with tempfile.TemporaryDirectory() as tmp:
            env = os.environ.copy()
            env["RELATIONSHIP_OS_STORE"] = "local"
            env["RELATIONSHIP_OS_LOCAL_DIR"] = tmp
            self.run_cli(["init", "--reset"], env)
            payload = {
                "contact_name": "Demo Client",
                "touch_date": "2026-05-20",
                "touchpoint_type": "coffee",
                "sentiment": "positive",
                "summary": "Coffee with Demo Client. Interested in retirement planning.",
                "raw_input": "Had coffee with Demo Client today. He is interested in retirement planning. Follow up next Friday.",
                "topics": ["retirement"],
                "action_items": "Follow up next Friday on retirement planning interest.",
                "contact_type": "prospect",
                "relationship_stage": "warm",
                "reminder_due": "2026-05-29",
                "reminder_priority": "medium",
                "reminder_context": "Follow up on retirement planning interest.",
            }
            out = self.run_cli(["log-touchpoint", "--json", json.dumps(payload)], env)
            self.assertIn("Logged coffee with Demo Client", out)
            self.assertIn("2026-05-29", out)

            contacts = (Path(tmp) / "Contacts.csv").read_text(encoding="utf-8")
            self.assertIn("Demo Client", contacts)

            prep = self.run_cli(["prep", "--name", "Demo Client"], env)
            self.assertIn("Prep: Demo Client", prep)
            self.assertIn("retirement", prep)

            today = self.run_cli(["today"], env)
            self.assertIn("Today:", today)

    def test_demo_command_populates_three_contacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            env = os.environ.copy()
            env["RELATIONSHIP_OS_STORE"] = "local"
            env["RELATIONSHIP_OS_LOCAL_DIR"] = tmp
            self.run_cli(["demo", "--reset"], env)
            contacts = (Path(tmp) / "Contacts.csv").read_text(encoding="utf-8")
            self.assertIn("Demo Client", contacts)
            self.assertIn("Sarah Lim", contacts)
            self.assertIn("Jason Wong", contacts)

    def test_sqlite_source_of_truth_syncs_csv_view(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)

            self.run_cli(["init", "--reset"], env)
            payload = self.sample_touchpoint_payload(
                reminder_due="2026-05-29",
                reminder_priority="medium",
                reminder_context="Follow up on retirement planning interest.",
            )
            out = self.run_cli(["log-touchpoint", "--json", json.dumps(payload)], env)

            self.assertIn("Logged coffee with Demo Client", out)
            self.assertTrue((base / "relationship_os.sqlite3").exists())
            contacts = (base / "local_sheet" / "Contacts.csv").read_text(encoding="utf-8")
            self.assertIn("Demo Client", contacts)
            reminders = (base / "local_sheet" / "Reminders.csv").read_text(encoding="utf-8")
            self.assertIn("2026-05-29", reminders)  # "next Friday" is no longer mis-parsed

            status = self.run_cli(["status"], env)
            self.assertIn("Relationship OS store: sqlite", status)
            self.assertIn("Consultant views: csv", status)

    def test_update_setting_updates_sqlite_and_csv_view(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            updated = self.run_json_cli(
                ["update-setting", "--key", "design_scheme", "--value", "awm-dark"], env,
            )
            self.assertTrue(updated["ok"])
            self.assertEqual(updated["value"], "awm-dark")

            with sqlite3.connect(base / "relationship_os.sqlite3") as conn:
                row = conn.execute(
                    "SELECT value FROM settings WHERE key = 'design_scheme'"
                ).fetchone()
            self.assertEqual(row[0], "awm-dark")
            settings_csv = (base / "local_sheet" / "Settings.csv").read_text(encoding="utf-8")
            self.assertIn("design_scheme,awm-dark", settings_csv)

            proc = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--env",
                    str(ROOT / ".env.example"),
                    "update-setting",
                    "--key",
                    "design_scheme",
                    "--value",
                    "purple",
                ],
                cwd=ROOT,
                text=True,
                capture_output=True,
                env=env,
            )
            self.assertNotEqual(proc.returncode, 0)
            self.assertIn("design_scheme", proc.stderr)

    def test_reminder_lifecycle_commands_update_rows_emit_events_and_sync_views(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            db_path = base / "relationship_os.sqlite3"

            self.run_cli(["init", "--reset"], env)
            payload = self.sample_touchpoint_payload(
                reminder_due="2026-05-29",
                reminder_priority="medium",
                reminder_context="Follow up on retirement planning interest.",
            )
            logged = self.run_json_cli(["log-touchpoint", "--json", json.dumps(payload)], env)
            reminder_id = logged["reminder"]["id"]

            completed = self.run_json_cli(
                ["complete-reminder", "--reminder-id", reminder_id],
                env,
            )
            self.assertEqual(completed["reminder"]["status"], "done")
            reminder = self.sqlite_row(db_path, "reminders", reminder_id)
            self.assertEqual(reminder["status"], "done")
            self.assertTrue(reminder["completed_at"])
            completed_events = self.run_json_cli(
                ["events", "--kind", "reminder_completed"],
                env,
            )
            self.assertEqual(completed_events["count"], 1)
            self.assertEqual(completed_events["events"][0]["subject_id"], reminder_id)
            self.assertEqual(completed_events["events"][0]["source"], "complete-reminder")
            self.assertIn("done", (base / "local_sheet" / "Reminders.csv").read_text(encoding="utf-8"))

            snooze_logged = self.run_json_cli(
                [
                    "log-touchpoint",
                    "--json",
                    json.dumps(
                        self.sample_touchpoint_payload(
                            contact_name="Snooze Client",
                            reminder_due="2026-05-30",
                            reminder_priority="high",
                            reminder_context="Send CPF follow-up.",
                        )
                    ),
                ],
                env,
            )
            snooze_id = snooze_logged["reminder"]["id"]
            snoozed = self.run_json_cli(
                [
                    "snooze-reminder",
                    "--reminder-id",
                    snooze_id,
                    "--new-due-date",
                    "2026-06-02",
                ],
                env,
            )
            self.assertEqual(snoozed["reminder"]["status"], "snoozed")
            reminder = self.sqlite_row(db_path, "reminders", snooze_id)
            self.assertEqual(reminder["status"], "snoozed")
            self.assertEqual(reminder["snoozed_until"], "2026-06-02")
            snoozed_events = self.run_json_cli(
                ["events", "--kind", "reminder_snoozed"],
                env,
            )
            self.assertEqual(snoozed_events["events"][0]["subject_id"], snooze_id)
            self.assertIn("2026-06-02", snoozed_events["events"][0]["payload"])

            cancel_logged = self.run_json_cli(
                [
                    "log-touchpoint",
                    "--json",
                    json.dumps(
                        self.sample_touchpoint_payload(
                            contact_name="Cancel Client",
                            reminder_due="2026-05-31",
                            reminder_priority="low",
                            reminder_context="No longer relevant.",
                        )
                    ),
                ],
                env,
            )
            cancel_id = cancel_logged["reminder"]["id"]
            cancelled = self.run_json_cli(
                ["cancel-reminder", "--reminder-id", cancel_id],
                env,
            )
            self.assertEqual(cancelled["reminder"]["status"], "cancelled")
            reminder = self.sqlite_row(db_path, "reminders", cancel_id)
            self.assertEqual(reminder["status"], "cancelled")
            self.assertTrue(reminder["completed_at"])
            cancelled_events = self.run_json_cli(
                ["events", "--kind", "reminder_cancelled"],
                env,
            )
            self.assertEqual(cancelled_events["events"][0]["subject_id"], cancel_id)
            self.assertEqual(cancelled_events["events"][0]["source"], "cancel-reminder")

    def test_events_query_filters_by_contact_kind_since_and_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            first = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload())],
                env,
            )
            second = self.run_json_cli(
                [
                    "log-touchpoint",
                    "--json",
                    json.dumps(self.sample_touchpoint_payload(contact_name="Sarah Lim")),
                ],
                env,
            )

            contact_events = self.run_json_cli(
                ["events", "--contact-id", first["contact"]["id"]],
                env,
            )
            self.assertGreaterEqual(contact_events["count"], 2)
            self.assertTrue(
                all(e["contact_id"] == first["contact"]["id"] for e in contact_events["events"])
            )

            touchpoint_events = self.run_json_cli(
                ["events", "--kind", "touchpoint_logged", "--limit", "1"],
                env,
            )
            self.assertEqual(touchpoint_events["count"], 1)
            self.assertEqual(touchpoint_events["events"][0]["kind"], "touchpoint_logged")
            self.assertIn(
                touchpoint_events["events"][0]["contact_id"],
                {first["contact"]["id"], second["contact"]["id"]},
            )

            future_events = self.run_json_cli(["events", "--since", "2099-01-01"], env)
            self.assertEqual(future_events["count"], 0)

    def test_sqlite_notes_columns_round_trip_to_csv_view(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            db_path = base / "relationship_os.sqlite3"
            self.run_cli(["init", "--reset"], env)

            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload())],
                env,
            )
            contact_id = logged["contact"]["id"]
            touchpoint_id = logged["touchpoint"]["id"]

            with sqlite3.connect(db_path) as conn:
                conn.execute(
                    'UPDATE contacts SET notes = ? WHERE id = ?',
                    ("Private contact note from app", contact_id),
                )
                conn.execute(
                    'UPDATE touchpoints SET notes = ? WHERE id = ?',
                    ("Private touchpoint annotation from app", touchpoint_id),
                )

            self.run_cli(["init"], env)

            contacts = (base / "local_sheet" / "Contacts.csv").read_text(encoding="utf-8")
            touchpoints = (base / "local_sheet" / "Touchpoints.csv").read_text(encoding="utf-8")
            self.assertIn("Private contact note from app", contacts)
            self.assertIn("Private touchpoint annotation from app", touchpoints)

    def test_log_touchpoint_validates_enums(self):
        rel = load_relationship_os_module()
        from datetime import date

        # Bad sentiment value should be rejected before any write.
        with self.assertRaises(rel.RelationshipOSError):
            rel.touchpoint_input_from_dict(
                {
                    "contact_name": "Demo",
                    "touchpoint_type": "coffee",
                    "sentiment": "happy",  # invalid
                    "summary": "x",
                },
                date(2026, 5, 20),
            )

        # Missing required field.
        with self.assertRaises(rel.RelationshipOSError):
            rel.touchpoint_input_from_dict(
                {"contact_name": "Demo", "touchpoint_type": "coffee", "sentiment": "positive"},
                date(2026, 5, 20),
            )

        # Bad reminder_due date format.
        with self.assertRaises(rel.RelationshipOSError):
            rel.touchpoint_input_from_dict(
                {
                    "contact_name": "Demo",
                    "touchpoint_type": "coffee",
                    "sentiment": "positive",
                    "summary": "x",
                    "reminder_due": "next Friday",  # not ISO
                },
                date(2026, 5, 20),
            )

        # Well-formed payload should produce a TouchpointInput.
        data = rel.touchpoint_input_from_dict(
            {
                "contact_name": "Demo Client",
                "touchpoint_type": "COFFEE",  # case-insensitive
                "sentiment": "Positive",
                "summary": "Coffee chat.",
                "topics": "retirement, protection",
                "reminder_due": "2026-05-29",
            },
            date(2026, 5, 20),
        )
        self.assertEqual(data.touchpoint_type, "coffee")
        self.assertEqual(data.sentiment, "positive")
        self.assertEqual(data.topics, ["retirement", "protection"])
        self.assertEqual(data.reminder_due, date(2026, 5, 29))

    def test_markdown_helpers_render_safe_frontmatter(self):
        rel = load_relationship_os_module()
        self.assertEqual(rel.safe_filename("A/B\\C: Demo Client?"), "a-b-c-demo-client")

        markdown = rel.render_frontmatter(
            {
                "id": "c_1",
                "type": "client",
                "tags": ["relationship-os", "client"],
            }
        ) + "Body"
        parsed = rel.parse_markdown_frontmatter(markdown)
        self.assertEqual(parsed["id"], "c_1")
        self.assertEqual(parsed["tags"], ["relationship-os", "client"])

    def test_design_token_parser_prefers_light_scheme(self):
        rel = load_relationship_os_module()
        with tempfile.TemporaryDirectory() as tmp:
            design = Path(tmp) / "design.md"
            design.write_text(
                """
:root[data-scheme="awm-dark"] {
  --ros-bg: #050505;
  --ros-text: #F5F1E8;
}
:root[data-scheme="awm-light"] {
  --ros-bg: #FAF7F0;
  --ros-paper: #FFFDF8;
  --ros-ink: #18231E;
  --ros-text: #1B1B18;
  --ros-text-muted: #777064;
  --ros-accent: #C6A34F;
}
Cormorant Garamond and Barlow are preferred.
""",
                encoding="utf-8",
            )
            tokens = rel.load_design_tokens(design, "awm-light")
            self.assertEqual(tokens["background"], "#FFFDF8")
            self.assertEqual(tokens["primary"], "#18231E")
            self.assertEqual(tokens["accent"], "#C6A34F")
            self.assertEqual(tokens["font_heading"], "Cormorant Garamond")
            self.assertEqual(tokens["font_body"], "Barlow")

    def test_brand_assets_from_design_are_inserted_into_ppt_and_pdf(self):
        rel = load_relationship_os_module()
        from PIL import Image
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            assets = tmp_path / "assets"
            assets.mkdir()
            for filename, color in {
                "logo_dark.png": (24, 35, 30, 255),
                "bg_mountain_clouds.png": (230, 235, 224, 255),
                "bg_network_mesh.png": (198, 163, 79, 255),
            }.items():
                Image.new("RGBA", (96, 48), color).save(assets / filename)

            design = tmp_path / "design.md"
            design.write_text(
                """
:root[data-scheme="awm-light"] {
  --ros-bg: #FAF7F0;
  --ros-paper: #FFFDF8;
  --ros-ink: #18231E;
  --ros-text: #1B1B18;
  --ros-text-muted: #777064;
  --ros-accent: #C6A34F;
}
Assets: logo_dark.png, bg_mountain_clouds.png, bg_network_mesh.png.
`Private FC operating aid. Not financial advice. Review before any client-facing use.`
""",
                encoding="utf-8",
            )
            original_assets_dir = os.environ.get("RELATIONSHIP_OS_BRAND_ASSETS_DIR")
            os.environ["RELATIONSHIP_OS_BRAND_ASSETS_DIR"] = str(assets)
            try:
                tokens = rel.load_design_tokens(design, "awm-light")
                self.assertEqual(Path(tokens["logo_cover_path"]), assets / "logo_dark.png")

                contact = {"name": "Demo Client", "type": "prospect", "relationship_stage": "warm"}
                pptx_path = tmp_path / "branded.pptx"
                rel.write_slides(pptx_path, contact, "annual review", [], [], tokens, max_slides=2)
                deck = Presentation(pptx_path)
                picture_count = sum(
                    1 for slide in deck.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                )
                self.assertGreaterEqual(picture_count, 2)

                pdf_path = tmp_path / "branded.pdf"
                rel.write_basic_pdf(pdf_path, "Appointment Summary: Demo Client", "# Appointment Summary\n\n## Notes\n- Review.", tokens)
                pdf_bytes = pdf_path.read_bytes()
                self.assertIn(b"/Subtype /Image", pdf_bytes)
                self.assertIn(b"Private FC operating aid", pdf_bytes)
            finally:
                if original_assets_dir is None:
                    os.environ.pop("RELATIONSHIP_OS_BRAND_ASSETS_DIR", None)
                else:
                    os.environ["RELATIONSHIP_OS_BRAND_ASSETS_DIR"] = original_assets_dir

    def test_markdown_export_is_idempotent_and_wikilinked(self):
        rel = load_relationship_os_module()
        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as vault_tmp:
            env = os.environ.copy()
            env["RELATIONSHIP_OS_STORE"] = "local"
            env["RELATIONSHIP_OS_LOCAL_DIR"] = tmp
            vault = Path(vault_tmp)

            self.run_cli(["demo", "--reset"], env)
            self.run_cli(["export-md", "--output-dir", str(vault)], env)

            contact_note = vault / "Contacts" / "demo-client.md"
            touchpoint_note = next((vault / "Touchpoints").glob("*demo-client*.md"))
            self.assertTrue(contact_note.exists())
            self.assertIn("[[Contacts/demo-client|Demo Client]]", touchpoint_note.read_text(encoding="utf-8"))

            original = contact_note.read_text(encoding="utf-8")
            contact_note.write_text(original + "\n## Consultant Notes\nKeep this manual note.\n", encoding="utf-8")
            self.run_cli(["export-md", "--output-dir", str(vault)], env)

            updated = contact_note.read_text(encoding="utf-8")
            self.assertEqual(updated.count(rel.AI_SECTION_START), 1)
            self.assertIn("Keep this manual note.", updated)

            index = rel.read_vault_index(vault)
            self.assertEqual(len(index["Contacts"]), 3)
            self.assertEqual(index["Contacts"][0]["type"], "prospect")

    def test_deliverable_generation_flows(self):
        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as vault_tmp:
            env = os.environ.copy()
            env["RELATIONSHIP_OS_STORE"] = "local"
            env["RELATIONSHIP_OS_LOCAL_DIR"] = tmp
            vault = Path(vault_tmp)

            self.run_cli(["demo", "--reset"], env)
            self.run_cli(["appointment-summary", "Demo", "Client", "--output-dir", str(vault)], env)
            self.run_cli(["proposal", "Demo", "Client", "retirement", "planning", "--output-dir", str(vault)], env)
            self.run_cli(["slides", "Demo", "Client", "annual", "review", "--output-dir", str(vault)], env)
            self.run_cli(["writeup", "retirement", "--output-dir", str(vault)], env)

            self.assertTrue(list((vault / "Generated" / "appointment_summary").glob("*.md")))
            self.assertTrue(list((vault / "Generated" / "appointment_summary").glob("*.pdf")))
            self.assertTrue(list((vault / "Generated" / "proposal").glob("*.md")))
            self.assertTrue(list((vault / "Generated" / "proposal").glob("*.pdf")))
            self.assertTrue(list((vault / "Generated" / "slides").glob("*.pptx")))
            self.assertTrue(list((vault / "Generated" / "writeup").glob("*.md")))

            dry_vault = vault / "dry"
            out = self.run_cli(
                ["proposal", "Demo", "Client", "retirement", "--output-dir", str(dry_vault), "--dry-run"],
                env,
            )
            self.assertIn("Would generate", out)
            out = self.run_cli(
                ["appointment-summary", "Demo", "Client", "--output-dir", str(dry_vault), "--dry-run"],
                env,
            )
            self.assertIn("Would generate", out)
            out = self.run_cli(
                ["slides", "Demo", "Client", "annual", "review", "--output-dir", str(dry_vault), "--dry-run"],
                env,
            )
            self.assertIn("Would generate", out)
            self.assertFalse(dry_vault.exists())

            dated = self.run_json_cli(
                [
                    "appointment-summary",
                    "Demo",
                    "Client",
                    "2026-05-24",
                    "--output-dir",
                    str(dry_vault),
                    "--dry-run",
                ],
                env,
            )
            self.assertTrue(
                any("2026-05-24 demo-client.md" in path for path in dated["files"])
            )

    def test_slide_section_validation(self):
        rel = load_relationship_os_module()

        # 'cover' and 'closing' must not be specified by Hermes (auto-added).
        with self.assertRaises(rel.RelationshipOSError):
            rel.slide_section_from_dict({"type": "cover"})
        with self.assertRaises(rel.RelationshipOSError):
            rel.slide_section_from_dict({"type": "closing"})

        # Unknown section type is rejected.
        with self.assertRaises(rel.RelationshipOSError):
            rel.slide_section_from_dict({"type": "feature_grid"})

        # content section requires both title and non-empty bullets.
        with self.assertRaises(rel.RelationshipOSError):
            rel.slide_section_from_dict({"type": "content", "bullets": ["a"]})
        with self.assertRaises(rel.RelationshipOSError):
            rel.slide_section_from_dict({"type": "content", "title": "X", "bullets": []})

        # touchpoint_detail requires both touchpoint_id and bullets.
        with self.assertRaises(rel.RelationshipOSError):
            rel.slide_section_from_dict({"type": "touchpoint_detail", "bullets": ["a"]})
        with self.assertRaises(rel.RelationshipOSError):
            rel.slide_section_from_dict({"type": "touchpoint_detail", "touchpoint_id": "t_1"})

        # Valid forms coerce correctly.
        section = rel.slide_section_from_dict(
            {"type": "content", "title": "Credit Cards 101", "bullets": ["Use UOB Lady's", "Stack cashback"]}
        )
        self.assertEqual(section.type, "content")
        self.assertEqual(section.title, "Credit Cards 101")
        self.assertEqual(section.bullets, ["Use UOB Lady's", "Stack cashback"])

        section = rel.slide_section_from_dict({"type": "touchpoint_summary", "limit": 3})
        self.assertEqual(section.limit, 3)

        section = rel.slide_section_from_dict({"type": "contact_context"})
        self.assertEqual(section.type, "contact_context")

    def test_content_section_paginates_at_six_bullets(self):
        rel = load_relationship_os_module()
        section = rel.SlideSection(
            type="content",
            title="Long Section",
            bullets=[f"bullet {i}" for i in range(1, 14)],  # 13 bullets
        )
        slides = rel.render_content_section(section)
        # 13 bullets at 6 per slide => 3 slides (6 + 6 + 1)
        self.assertEqual(len(slides), 3)
        self.assertEqual(slides[0].title, "Long Section")
        self.assertEqual(slides[1].title, "Long Section (cont.)")
        self.assertEqual(slides[2].title, "Long Section (cont.)")
        self.assertEqual(len(slides[0].bullets), 6)
        self.assertEqual(len(slides[1].bullets), 6)
        self.assertEqual(len(slides[2].bullets), 1)

    def test_build_deck_always_brackets_with_cover_and_closing(self):
        rel = load_relationship_os_module()
        contact = {"id": "c_1", "name": "Demo Client", "type": "prospect", "relationship_stage": "warm"}
        sections = [rel.SlideSection(type="contact_context")]
        deck = rel.build_deck(contact, "Annual review", sections, [], [])
        # cover + 1 content_context slide + closing = 3
        self.assertEqual(len(deck), 3)
        self.assertEqual(deck[0].type, "cover")
        self.assertEqual(deck[-1].type, "closing")
        self.assertIn("Demo Client", deck[0].title)

    def test_draft_notice_appears_in_every_deliverable(self):
        """Compliance guard: every generated MD/PDF/PPTX must contain
        the draft-only notice. A future agent editing the templates and
        accidentally removing the notice will fail this test."""
        rel = load_relationship_os_module()
        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as vault_tmp:
            env = os.environ.copy()
            env["RELATIONSHIP_OS_STORE"] = "local"
            env["RELATIONSHIP_OS_LOCAL_DIR"] = tmp
            vault = Path(vault_tmp)

            self.run_cli(["demo", "--reset"], env)
            self.run_cli(["appointment-summary", "Demo", "Client", "--output-dir", str(vault)], env)
            self.run_cli(["proposal", "Demo", "Client", "retirement", "planning", "--output-dir", str(vault)], env)
            self.run_cli(["slides", "Demo", "Client", "annual", "review", "--output-dir", str(vault)], env)
            self.run_cli(["writeup", "retirement", "--output-dir", str(vault)], env)

            draft_phrase = "Draft only"

            # All MD outputs.
            for md in (vault / "Generated").glob("**/*.md"):
                text = md.read_text(encoding="utf-8")
                self.assertIn(
                    draft_phrase,
                    text,
                    f"DRAFT_NOTICE missing from {md.relative_to(vault)}",
                )

            # All PDF outputs (search the raw bytes; pdf_escape never alters
            # ASCII "Draft only").
            for pdf in (vault / "Generated").glob("**/*.pdf"):
                data = pdf.read_bytes()
                self.assertIn(
                    draft_phrase.encode("ascii"),
                    data,
                    f"DRAFT_NOTICE missing from {pdf.relative_to(vault)}",
                )

            # All PPTX outputs.
            from pptx import Presentation
            for pptx in (vault / "Generated").glob("**/*.pptx"):
                deck = Presentation(pptx)
                full_text = []
                for slide in deck.slides:
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                for run in paragraph.runs:
                                    full_text.append(run.text)
                self.assertIn(
                    draft_phrase,
                    " ".join(full_text),
                    f"DRAFT_NOTICE missing from {pptx.relative_to(vault)}",
                )

    def test_slides_json_payload_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as vault_tmp:
            env = os.environ.copy()
            env["RELATIONSHIP_OS_STORE"] = "local"
            env["RELATIONSHIP_OS_LOCAL_DIR"] = tmp
            vault = Path(vault_tmp)
            self.run_cli(["demo", "--reset"], env)

            payload = {
                "contact_name": "Demo Client",
                "purpose": "Annual review + credit cards primer",
                "sections": [
                    {"type": "contact_context"},
                    {"type": "touchpoint_summary"},
                    {"type": "content",
                     "title": "Maximising Credit Cards in Singapore",
                     "bullets": [
                         "Optimise for highest-spend categories per card",
                         "Stack with bank-specific cashback events",
                         "Track expiry of welcome bonuses",
                         "Request annual fee waivers proactively",
                     ]},
                    {"type": "open_items"},
                ],
            }
            out = self.run_cli(
                ["slides", "--json", json.dumps(payload), "--output-dir", str(vault)],
                env,
            )
            self.assertIn("Generated", out)
            decks = list((vault / "Generated" / "slides").glob("*.pptx"))
            self.assertEqual(len(decks), 1)

            # Verify cover + sections + closing landed.
            from pptx import Presentation
            deck = Presentation(decks[0])
            # 1 cover + contact_context (1) + touchpoint_summary (1) + content (1) + open_items (1) + closing = 6
            self.assertEqual(len(deck.slides), 6)

    # ---- Pass A additions: dedup, find-reminders, merge-contacts, update-contact ----

    def test_log_touchpoint_skips_duplicate_reminder(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            # First log: creates Demo Client + reminder for 2026-05-29.
            first = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        reminder_due="2026-05-29",
                        reminder_priority="medium",
                        reminder_context="Follow up on retirement planning interest.",
                    )
                )],
                env,
            )
            self.assertIn("reminder", first)
            existing_reminder_id = first["reminder"]["id"]

            # Second log: same contact, slightly different due date, near-identical context.
            # Should skip the new reminder and emit reminder_duplicate_skipped.
            second = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        touch_date="2026-05-21",
                        reminder_due="2026-05-30",  # within ±7 days of 5-29
                        reminder_priority="medium",
                        reminder_context="Follow up about retirement planning.",
                    )
                )],
                env,
            )
            self.assertNotIn("reminder", second)
            self.assertIn("reminder_skipped", second)
            self.assertEqual(second["reminder_skipped"]["existing_reminder_id"], existing_reminder_id)

            # Events table should record the skip.
            skip_events = self.run_json_cli(
                ["events", "--kind", "reminder_duplicate_skipped"], env
            )
            self.assertEqual(skip_events["count"], 1)

            # Reminders table should still have exactly 1 reminder for Demo Client.
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                rows = conn.execute(
                    "SELECT id FROM reminders WHERE status = 'pending'"
                ).fetchall()
            self.assertEqual(len(rows), 1)

    def test_log_touchpoint_allows_duplicate_with_flag(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            payload = self.sample_touchpoint_payload(
                reminder_due="2026-05-29",
                reminder_priority="medium",
                reminder_context="Follow up on retirement planning interest.",
            )
            self.run_json_cli(["log-touchpoint", "--json", json.dumps(payload)], env)

            payload2 = dict(payload)
            payload2["allow_duplicate_reminder"] = True
            second = self.run_json_cli(["log-touchpoint", "--json", json.dumps(payload2)], env)
            self.assertIn("reminder", second)
            self.assertNotIn("reminder_skipped", second)

            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                rows = conn.execute("SELECT COUNT(*) FROM reminders").fetchone()
            self.assertEqual(rows[0], 2)

    def test_log_touchpoint_distinct_reminders_not_deduped(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        reminder_due="2026-05-29",
                        reminder_context="Follow up on retirement planning.",
                    )
                )],
                env,
            )
            # Same contact, different due date (>7 days away) AND different context.
            second = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        reminder_due="2026-08-15",
                        reminder_context="Send birthday greeting in August.",
                    )
                )],
                env,
            )
            self.assertIn("reminder", second)
            self.assertNotIn("reminder_skipped", second)

    def test_find_reminders_matches_context_substring(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        reminder_due="2026-06-05",
                        reminder_context="Send SRS illustration this week.",
                    )
                )],
                env,
            )
            self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        contact_name="Sarah Lim",
                        reminder_due="2026-06-10",
                        reminder_context="Discuss protection planning options.",
                    )
                )],
                env,
            )

            srs = self.run_json_cli(["find-reminders", "--query", "SRS"], env)
            self.assertEqual(srs["count"], 1)
            self.assertIn("SRS", srs["candidates"][0]["context"])

            protection = self.run_json_cli(["find-reminders", "--query", "protection"], env)
            self.assertEqual(protection["count"], 1)
            self.assertEqual(protection["candidates"][0]["contact_name"], "Sarah Lim")

    def test_find_reminders_status_filter(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        reminder_due="2026-06-05",
                        reminder_context="Send SRS illustration.",
                    )
                )],
                env,
            )
            rid = logged["reminder"]["id"]
            # Complete it.
            self.run_json_cli(["complete-reminder", "--reminder-id", rid], env)

            pending = self.run_json_cli(["find-reminders", "--query", "SRS"], env)
            self.assertEqual(pending["count"], 0)
            done = self.run_json_cli(["find-reminders", "--query", "SRS", "--status", "done"], env)
            self.assertEqual(done["count"], 1)
            any_status = self.run_json_cli(["find-reminders", "--query", "SRS", "--status", "any"], env)
            self.assertEqual(any_status["count"], 1)

    def test_find_reminders_empty_query_returns_all_pending(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            for due, ctx in [("2026-06-01", "first"), ("2026-06-02", "second"), ("2026-06-03", "third")]:
                self.run_json_cli(
                    ["log-touchpoint", "--json", json.dumps(
                        self.sample_touchpoint_payload(
                            contact_name=f"Contact {ctx}",
                            reminder_due=due,
                            reminder_context=f"Reminder for {ctx} interaction.",
                        )
                    )],
                    env,
                )
            out = self.run_json_cli(["find-reminders"], env)
            self.assertEqual(out["count"], 3)

    def test_merge_contacts_moves_touchpoints_and_reminders(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            # Create two contacts via log-touchpoint.
            a = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        contact_name="Sarah",
                        reminder_due="2026-06-01",
                        reminder_context="Follow up Sarah.",
                    )
                )],
                env,
            )
            b = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        contact_name="Sarah Lim",
                        reminder_due="2026-07-01",  # different month to avoid dedup
                        reminder_context="Discuss something completely different.",
                    )
                )],
                env,
            )
            from_id = a["contact"]["id"]
            into_id = b["contact"]["id"]

            merged = self.run_json_cli(
                ["merge-contacts", "--from", from_id, "--into", into_id], env
            )
            self.assertGreaterEqual(merged["touchpoints_moved"], 1)
            self.assertGreaterEqual(merged["reminders_moved"], 1)

            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                contacts = conn.execute("SELECT id, name FROM contacts").fetchall()
                touchpoints_for_into = conn.execute(
                    "SELECT COUNT(*) AS n FROM touchpoints WHERE contact_id = ?", (into_id,)
                ).fetchone()["n"]
                reminders_for_into = conn.execute(
                    "SELECT COUNT(*) AS n FROM reminders WHERE contact_id = ?", (into_id,)
                ).fetchone()["n"]

            ids = [row["id"] for row in contacts]
            self.assertIn(into_id, ids)
            self.assertNotIn(from_id, ids)
            self.assertGreaterEqual(touchpoints_for_into, 2)
            self.assertGreaterEqual(reminders_for_into, 2)

            events = self.run_json_cli(["events", "--kind", "contact_merged"], env)
            self.assertEqual(events["count"], 1)

    def test_merge_contacts_field_strategy(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)

            # Create two contacts via log-touchpoint.
            a = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        contact_name="Vincent Old",
                        reminder_due="2026-06-01",
                        reminder_context="Random.",
                    )
                )],
                env,
            )
            b = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        contact_name="Vincent Canonical",
                        reminder_due="2026-07-01",
                        reminder_context="Other.",
                    )
                )],
                env,
            )
            # Populate fields on both contacts.
            self.run_json_cli(["update-contact", "--id", a["contact"]["id"],
                               "--field", "phone", "--value", "+65 9111 1111"], env)
            self.run_json_cli(["update-contact", "--id", a["contact"]["id"],
                               "--field", "family", "--value", "Wife and one kid."], env)
            self.run_json_cli(["update-contact", "--id", b["contact"]["id"],
                               "--field", "family", "--value", "Wife and one toddler."], env)

            merged = self.run_json_cli(
                ["merge-contacts", "--from", a["contact"]["id"],
                 "--into", b["contact"]["id"]], env,
            )
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    "SELECT phone, family, notes FROM contacts WHERE id = ?",
                    (b["contact"]["id"],),
                ).fetchone()
            # Phone was empty on into; copied from from.
            self.assertEqual(row["phone"], "+65 9111 1111")
            # Family was populated on both; into wins, from's value lands in notes.
            self.assertIn("toddler", row["family"])
            self.assertIn("merged from Vincent Old", row["notes"])
            self.assertIn("kid", row["notes"])

    def test_merge_contacts_rejects_unknown_id(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            proc = subprocess.run(
                [sys.executable, str(SCRIPT), "--env", str(ROOT / ".env.example"),
                 "merge-contacts", "--from", "c_does_not_exist", "--into", "c_also_no"],
                cwd=ROOT, text=True, capture_output=True, env=env,
            )
            self.assertNotEqual(proc.returncode, 0)
            self.assertIn("No contact with id", proc.stderr)

    def test_update_contact_append_field_prepends_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload())], env,
            )
            cid = logged["contact"]["id"]
            self.run_json_cli(["update-contact", "--id", cid,
                               "--field", "family", "--value", "Wife and one kid."], env)
            self.run_json_cli(["update-contact", "--id", cid,
                               "--field", "family", "--value", "Wife pregnant with second."], env)
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    "SELECT family FROM contacts WHERE id = ?", (cid,)
                ).fetchone()
            family = row["family"]
            self.assertIn("Wife and one kid", family)
            self.assertIn("Wife pregnant with second", family)
            # Two date-prefixed entries.
            self.assertEqual(family.count("["), 2)

    def test_update_contact_replace_field_overwrites(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload())], env,
            )
            cid = logged["contact"]["id"]
            self.run_json_cli(["update-contact", "--id", cid,
                               "--field", "phone", "--value", "+65 9111 1111"], env)
            self.run_json_cli(["update-contact", "--id", cid,
                               "--field", "phone", "--value", "+65 9222 2222"], env)
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    "SELECT phone FROM contacts WHERE id = ?", (cid,)
                ).fetchone()
            self.assertEqual(row["phone"], "+65 9222 2222")

    def test_update_contact_rejects_hermes_managed_field(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload())], env,
            )
            cid = logged["contact"]["id"]
            proc = subprocess.run(
                [sys.executable, str(SCRIPT), "--env", str(ROOT / ".env.example"),
                 "update-contact", "--id", cid, "--field", "relationship_stage", "--value", "hot"],
                cwd=ROOT, text=True, capture_output=True, env=env,
            )
            self.assertNotEqual(proc.returncode, 0)
            self.assertIn("relationship_stage", proc.stderr)

    def test_update_contact_emits_diff_event(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload())], env,
            )
            cid = logged["contact"]["id"]
            self.run_json_cli(["update-contact", "--id", cid,
                               "--field", "occupation", "--value", "software developer"], env)
            events = self.run_json_cli(
                ["events", "--kind", "contact_updated", "--contact-id", cid], env,
            )
            self.assertGreaterEqual(events["count"], 1)
            latest = events["events"][0]
            self.assertEqual(latest["source"], "hermes:update-contact-fields")
            payload = json.loads(latest["payload"])
            self.assertIn("occupation", payload["diff"])

    def test_today_brief_returns_structured_composite(self):
        """today-brief returns reminders_due + birthdays_today + ripe_signals
        + debt_top + stats, plus a Telegram-ready brief_text."""
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            # Log a touchpoint with a reminder due today.
            payload = self.sample_touchpoint_payload(contact_name="Vince")
            payload["reminder_due"] = "2026-05-25"  # arbitrary "today" for the test
            payload["reminder_priority"] = "high"
            payload["reminder_context"] = "Send the SRS sheet"
            payload["reminder_type"] = "follow_up"
            self.run_json_cli(["log-touchpoint", "--json", json.dumps(payload)], env)
            brief = self.run_json_cli(["today-brief"], env)
            self.assertEqual(brief.get("command"), "today-brief")
            self.assertIn("reminders_due", brief)
            self.assertIn("birthdays_today", brief)
            self.assertIn("ripe_signals", brief)
            self.assertIn("debt_top", brief)
            self.assertIn("stats", brief)
            self.assertIn("brief_text", brief)
            # Stats keys
            stats = brief["stats"]
            self.assertIn("reminders_due_count", stats)
            self.assertIn("birthdays_today_count", stats)
            self.assertIn("ripe_count", stats)
            self.assertIn("debt_total_count", stats)
            # brief_text is non-empty and references the consultant name greeting
            self.assertIn("Good morning", brief["brief_text"])

    def test_today_brief_text_format_returns_human_readable(self):
        """When run with --format=text, today-brief outputs the brief_text
        directly (Telegram-ready, no JSON)."""
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            text = self.run_cli(["today-brief"], env)  # default format is text
            self.assertIn("Good morning", text)
            self.assertIn("Reminders due today", text)
            self.assertIn("Birthdays today", text)
            self.assertIn("Ripe to reach out", text)
            self.assertIn("Relationship debt", text)
            # No raw JSON should leak when format=text
            self.assertNotIn('"ok":', text)
            self.assertNotIn('"command":', text)

    def test_update_contact_respects_source_override_from_payload(self):
        """JSON payload can override the event source (used by the app when
        shelling out, so Settings' "last Hermes event" filter excludes
        app-initiated edits via the source NOT LIKE 'app:%' guard).
        """
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload())], env,
            )
            cid = logged["contact"]["id"]
            payload = {
                "id": cid,
                "updates": {"occupation": "software developer"},
                "source": "app:edit-contact",
            }
            self.run_json_cli(["update-contact", "--json", json.dumps(payload)], env)
            events = self.run_json_cli(
                ["events", "--kind", "contact_updated", "--contact-id", cid], env,
            )
            self.assertGreaterEqual(events["count"], 1)
            self.assertEqual(events["events"][0]["source"], "app:edit-contact")

    def test_update_contact_by_name_resolves_unique(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload(contact_name="Vincent"))], env,
            )
            updated = self.run_json_cli(
                ["update-contact", "--name", "Vincent",
                 "--field", "occupation", "--value", "software developer"], env,
            )
            self.assertEqual(updated["contact_name"], "Vincent")
            self.assertIn("occupation", updated["diff"])

    # ---- Pass B additions: policies module ----

    def _seed_contact(self, env: Dict[str, str], name: str = "Demo Client") -> str:
        logged = self.run_json_cli(
            ["log-touchpoint", "--json", json.dumps(self.sample_touchpoint_payload(contact_name=name))], env,
        )
        return logged["contact"]["id"]

    def test_add_policy_basic(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            added = self.run_json_cli(
                [
                    "add-policy",
                    "--contact-id", cid,
                    "--insurer", "AIA",
                    "--plan-name", "Secure Flexi Term",
                    "--policy-type", "protection",
                    "--premium-amount", "$2,050",
                    "--premium-frequency", "annual",
                ],
                env,
            )
            self.assertTrue(added["ok"])
            self.assertEqual(added["policy"]["insurer"], "AIA")
            self.assertEqual(added["policy"]["plan_name"], "Secure Flexi Term")
            self.assertEqual(added["policy"]["status"], "active")

            events = self.run_json_cli(["events", "--kind", "policy_created"], env)
            self.assertEqual(events["count"], 1)

    def test_policy_extended_fields_round_trip_to_sqlite_csv_and_prep(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            payload = {
                "contact_id": cid,
                "insurer": "AIA",
                "plan_name": "Platinum Wealth Venture",
                "policy_type": "investment-linked",
                "policy_number": "PWV-123",
                "premium_amount": "S$24,000",
                "premium_frequency": "annual",
                "premium_term": "10 years",
                "policy_term": "whole life",
                "payment_method": "GIRO",
                "start_date": "2025-10-17",
                "review_date": "2026-10-17",
                "review_frequency": "annual",
                "last_reviewed": "2026-05-20",
                "current_value": "S$35,423",
                "valuation_date": "2026-05-22",
                "surrender_value": "S$30,000",
                "policy_owner": "Vincent's mother",
                "life_assured": "Vincent",
                "payor": "Vincent",
                "beneficiaries": "Nomination pending",
                "riders": "Early CI rider",
                "servicing_rep": "Lucas",
                "needs_category": "retirement",
            }
            added = self.run_json_cli(["add-policy", "--json", json.dumps(payload)], env)
            pid = added["policy"]["id"]
            self.assertEqual(added["policy"]["current_value"], "S$35,423")
            self.assertEqual(added["policy"]["policy_owner"], "Vincent's mother")
            self.assertEqual(added["policy"]["life_assured"], "Vincent")

            row = self.sqlite_row(base / "relationship_os.sqlite3", "policies", pid)
            self.assertEqual(row["policy_number"], "PWV-123")
            self.assertEqual(row["beneficiaries"], "Nomination pending")

            policies_csv = (base / "local_sheet" / "Policies.csv").read_text(encoding="utf-8")
            self.assertIn("current_value", policies_csv.splitlines()[0])
            self.assertIn("S$35,423", policies_csv)

            prep = self.run_cli(["prep", "--name", "Vincent"], env)
            self.assertIn("value S$35,423 as of 2026-05-22", prep)
            self.assertIn("owner Vincent's mother", prep)

    def test_add_policy_requires_insurer_and_plan(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            for missing in [
                ["add-policy", "--contact-id", cid, "--plan-name", "Plan"],
                ["add-policy", "--contact-id", cid, "--insurer", "AIA"],
            ]:
                proc = subprocess.run(
                    [sys.executable, str(SCRIPT), "--env", str(ROOT / ".env.example"), *missing],
                    cwd=ROOT, text=True, capture_output=True, env=env,
                )
                self.assertNotEqual(proc.returncode, 0)

    def test_add_policy_rejects_bad_status(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            proc = subprocess.run(
                [sys.executable, str(SCRIPT), "--env", str(ROOT / ".env.example"),
                 "add-policy", "--contact-id", cid, "--insurer", "AIA",
                 "--plan-name", "X", "--status", "fictional"],
                cwd=ROOT, text=True, capture_output=True, env=env,
            )
            self.assertNotEqual(proc.returncode, 0)
            self.assertIn("status", proc.stderr.lower())

    def test_update_policy_changes_fields_and_emits_event(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            added = self.run_json_cli(
                ["add-policy", "--contact-id", cid, "--insurer", "AIA",
                 "--plan-name", "Secure Flexi Term", "--premium-amount", "$2,050"], env,
            )
            pid = added["policy"]["id"]
            updated = self.run_json_cli(
                ["update-policy", "--id", pid,
                 "--premium-amount", "$2,400", "--policy-type", "protection",
                 "--current-value", "S$12,000", "--valuation-date", "2026-05-22"], env,
            )
            self.assertIn("premium_amount", updated["diff"])
            self.assertIn("policy_type", updated["diff"])
            self.assertIn("current_value", updated["diff"])
            self.assertEqual(updated["policy"]["valuation_date"], "2026-05-22")
            events = self.run_json_cli(["events", "--kind", "policy_updated"], env)
            self.assertEqual(events["count"], 1)

    def test_update_policy_rejects_unknown_id(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            proc = subprocess.run(
                [sys.executable, str(SCRIPT), "--env", str(ROOT / ".env.example"),
                 "update-policy", "--id", "p_does_not_exist", "--insurer", "X"],
                cwd=ROOT, text=True, capture_output=True, env=env,
            )
            self.assertNotEqual(proc.returncode, 0)

    def test_list_policies_filters_archived_by_default(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            a = self.run_json_cli(
                ["add-policy", "--contact-id", cid, "--insurer", "AIA", "--plan-name", "Active Plan"], env,
            )
            b = self.run_json_cli(
                ["add-policy", "--contact-id", cid, "--insurer", "GE", "--plan-name", "To Archive"], env,
            )
            self.run_json_cli(["archive-policy", "--id", b["policy"]["id"]], env)

            default = self.run_json_cli(["list-policies", "--contact-id", cid], env)
            self.assertEqual(default["count"], 1)
            self.assertEqual(default["policies"][0]["id"], a["policy"]["id"])

            include = self.run_json_cli(
                ["list-policies", "--contact-id", cid, "--include-archived"], env,
            )
            self.assertEqual(include["count"], 2)

    def test_archive_policy_sets_status_not_deletes(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            added = self.run_json_cli(
                ["add-policy", "--contact-id", cid, "--insurer", "AIA", "--plan-name", "X"], env,
            )
            pid = added["policy"]["id"]
            self.run_json_cli(["archive-policy", "--id", pid], env)

            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    "SELECT id, status FROM policies WHERE id = ?", (pid,),
                ).fetchone()
            self.assertIsNotNone(row)
            self.assertEqual(row["status"], "archived")

            events = self.run_json_cli(["events", "--kind", "policy_archived"], env)
            self.assertEqual(events["count"], 1)

    # ---- Pass D additions: archive/rename contact + dedup behavior ----

    def test_archive_contact_sets_archived_at_and_emits_event(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Hayden")
            self.run_json_cli(["archive-contact", "--id", cid], env)
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    "SELECT archived_at FROM contacts WHERE id = ?", (cid,)
                ).fetchone()
            self.assertTrue(row["archived_at"])
            events = self.run_json_cli(["events", "--kind", "contact_archived"], env)
            self.assertEqual(events["count"], 1)

    def test_unarchive_contact_clears_archived_at(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Hayden")
            self.run_json_cli(["archive-contact", "--id", cid], env)
            self.run_json_cli(["unarchive-contact", "--id", cid], env)
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    "SELECT archived_at FROM contacts WHERE id = ?", (cid,)
                ).fetchone()
            self.assertFalse(row["archived_at"])

    def test_rename_contact_cascades_to_touchpoints_and_reminders(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            logged = self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(
                        contact_name="Hayden",
                        reminder_due="2026-06-15",
                        reminder_context="A reminder.",
                    )
                )],
                env,
            )
            cid = logged["contact"]["id"]
            result = self.run_json_cli(
                ["rename-contact", "--id", cid, "--new-name", "Hayden Foo"], env,
            )
            self.assertEqual(result["to"], "Hayden Foo")
            self.assertEqual(result["touchpoints_updated"], 1)
            self.assertEqual(result["reminders_updated"], 1)
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                contact = conn.execute(
                    "SELECT name FROM contacts WHERE id = ?", (cid,)
                ).fetchone()
                tp = conn.execute(
                    "SELECT contact_name FROM touchpoints WHERE contact_id = ?", (cid,)
                ).fetchone()
                rem = conn.execute(
                    "SELECT contact_name FROM reminders WHERE contact_id = ?", (cid,)
                ).fetchone()
            self.assertEqual(contact["name"], "Hayden Foo")
            self.assertEqual(tp["contact_name"], "Hayden Foo")
            self.assertEqual(rem["contact_name"], "Hayden Foo")

    def test_rename_contact_rejects_existing_name_collision(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            self._seed_contact(env, "Hayden Foo")
            cid_b = self._seed_contact(env, "Hayden")
            proc = subprocess.run(
                [sys.executable, str(SCRIPT), "--env", str(ROOT / ".env.example"),
                 "rename-contact", "--id", cid_b, "--new-name", "Hayden Foo"],
                cwd=ROOT, text=True, capture_output=True, env=env,
            )
            self.assertNotEqual(proc.returncode, 0)
            self.assertIn("already exists", proc.stderr)

    def test_import_touchpoint_does_not_update_last_touch_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            # Log a real meeting first.
            self.run_json_cli(
                ["log-touchpoint", "--json", json.dumps(
                    self.sample_touchpoint_payload(touch_date="2026-05-15")
                )],
                env,
            )
            # Then an import touchpoint a week later — should NOT bump last_touch_date.
            payload = self.sample_touchpoint_payload(touch_date="2026-05-22")
            payload["touchpoint_type"] = "import"
            payload["summary"] = "Imported from clients.xlsx"
            payload["raw_input"] = "Imported from clients.xlsx"
            self.run_json_cli(["log-touchpoint", "--json", json.dumps(payload)], env)
            db_path = base / "relationship_os.sqlite3"
            with sqlite3.connect(db_path) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    "SELECT last_touch_date FROM contacts WHERE name = 'Demo Client'"
                ).fetchone()
            self.assertEqual(row["last_touch_date"], "2026-05-15")

    def test_prep_includes_active_policies(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            env = self.sqlite_csv_env(base)
            self.run_cli(["init", "--reset"], env)
            cid = self._seed_contact(env, "Vincent")
            self.run_json_cli(
                ["add-policy", "--contact-id", cid, "--insurer", "AIA",
                 "--plan-name", "Secure Flexi Term", "--premium-amount", "$2,050"], env,
            )
            prep = self.run_json_cli(["prep", "--name", "Vincent"], env)
            self.assertIn("policies", prep)
            self.assertEqual(len(prep["policies"]), 1)
            self.assertEqual(prep["policies"][0]["plan_name"], "Secure Flexi Term")


if __name__ == "__main__":
    unittest.main()
